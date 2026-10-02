"""Migração única da planilha Contabilidade.xlsx para lancamentos e recorrencias (Rodada 2, etapa R2.1).

    python -m app.migrar_financeiro [--gravar] [--arquivo CAMINHO]

Sem --gravar só simula e imprime o relatório. Com --gravar faz backup, depois INSERT OR IGNORE pela coluna
origem (aba:linha), então pode rodar de novo sem duplicar. A planilha é só lida, nunca salva.
Regras e armadilhas: docs/mapa-financeiro.md, seção 6.
"""
import argparse
import datetime as dt
import re
import sqlite3
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from . import db
from .financeiro import vencimento
from .importar import _txt

INICIO = "2026-10"   # recorrências começam no mês em que o CRM passa a valer
COLUNAS = ("escopo", "tipo", "data", "pago_em", "valor", "categoria", "descricao", "contraparte", "cliente_id",
           "conta", "forma", "parcela_num", "parcela_total", "nfse", "observacao", "origem")
COLUNAS_REC = ("escopo", "descricao", "categoria", "valor", "dia_vencimento", "fixo", "conta", "inicio", "fim", "origem")
CUSTOS = [  # (aba, tabela, escopo, fixo)
    ("Custos Empresariais", "tblCustosEmpresaFixos", "PJ", 1),
    ("Custos Empresariais", "tblCustosEmpresaVariaveis", "PJ", 0),
    ("Custos Pessoais", "tblCustosPessoaisFixos", "PF", 1),
    ("Custos Pessoais", "tblCustosPessoaisVariaveis", "PF", 0),
]


def _norm(v) -> str:
    """Sem acento, minúsculo, espaços únicos: para comparar nomes."""
    s = unicodedata.normalize("NFD", str(v or "")).encode("ascii", "ignore").decode()
    return " ".join(s.casefold().split())


def _linhas(wb, aba: str, tabela: str):
    """(linha da planilha, dict cabeçalho->valor). Igual a importar.ler_tabela, mas devolve o número da linha,
    que vira a origem. Pula linhas totalmente vazias."""
    ws = wb[aba]
    linhas = ws[ws.tables[tabela].ref]
    cab = [str(c.value).strip() for c in linhas[0]]
    for linha in linhas[1:]:
        vals = [c.value for c in linha]
        if any(v not in (None, "") for v in vals):
            yield linha[0].row, dict(zip(cab, vals))


def _cat(v) -> str:
    c = _txt(v)
    return "Outros" if c in (None, "Outro") else c


def _forma(v):
    f = _txt(v)
    return "Cartão Débito" if f == "Débito" else f


def _tipo(r):
    return {"entrada": "entrada", "saida": "saida"}.get(_norm(r.get("Tipo")))


def _pago(r, iso):
    return iso if _norm(r.get("Status")) == "pago" else None


def _conta_pf(descricao):
    """Mesma regra da planilha (linha do tempo U41): Mercado Pago se a descrição cita, senão Inter."""
    return "Mercado Pago" if "mercado pago" in _norm(descricao) else "Inter"


def _parcela(v):
    m = re.fullmatch(r"(\d+)\s*de\s*(\d+)", str(v or "").strip())
    return (int(m[1]), int(m[2])) if m else (None, None)


def _obs(*partes):
    return " | ".join(p for p in map(_txt, partes) if p) or None


def _percorrer(wb, p, rotulo, aba, tabela, prefixo, montar, col_data="Data", col_valor="Valor (R$)"):
    """Lê uma tabela, descarta o que não é lançamento (com motivo) e acumula os aproveitados em p['lancamentos']."""
    for linha, r in _linhas(wb, aba, tabela):
        p["lidas"][rotulo] += 1
        d = r.get(col_data)
        if not isinstance(d, dt.datetime):   # total, linha fantasma ou texto: só data de verdade vira lançamento
            motivo = "data em texto (linha de total)" if isinstance(d, str) else "sem data (total ou linha fantasma)"
            p["descartes"][f"{rotulo}: {motivo}"] += 1
            continue
        try:
            bruto = float(r.get(col_valor))
        except (TypeError, ValueError):
            bruto = 0.0
        valor = round(bruto, 2)
        if valor <= 0:
            p["descartes"][f"{rotulo}: valor ausente ou não positivo"] += 1
            continue
        iso = d.date().isoformat()
        extra = montar(r, iso)
        if extra is None:
            p["descartes"][f"{rotulo}: fora da regra (tipo inválido ou Incluir? = Não)"] += 1
            continue
        p["arredondados"] += abs(bruto - valor) > 1e-9
        p["lancamentos"].append(dict.fromkeys(COLUNAS) | {"data": iso, "valor": valor, "origem": f"{prefixo}:{linha}"} | extra)


def _custos(wb, p):
    """Custos fixos/variáveis viram recorrência; eventual vira 1 lançamento pendente. Se a obrigação já existe como
    pendente nas transações, não cria nada (senão a conta apareceria em dobro)."""
    pendentes = {(l["escopo"], _norm(x)) for l in p["lancamentos"] if not l["pago_em"] and l["origem"][:3] in ("PJ:", "PF:")
                 for x in (l["descricao"], l["contraparte"]) if x}
    for aba, tabela, escopo, fixo in CUSTOS:
        for linha, r in _linhas(wb, aba, tabela):
            p["lidas"]["Custos"] += 1
            desc = _txt(r.get("Descrição"))
            if not desc:
                p["descartes"]["Custos: linha vazia do modelo"] += 1
                continue
            try:
                valor, dia = round(float(r.get("Valor")), 2), int(r.get("Dia Venc."))
            except (TypeError, ValueError):
                valor, dia = 0, 0
            if valor <= 0 or not 1 <= dia <= 31:
                p["descartes"]["Custos: sem valor ou dia de vencimento"] += 1
                continue
            if _txt(r.get("Data pagamento")):   # armadilha 8: já está nas transações
                p["descartes"]["Custos: 'Data pagamento' não vira lançamento (já está nas transações)"] += 1
            obs = _norm(r.get("Observação"))
            if (escopo, _norm(desc)) in pendentes:
                p["descartes"][f"Custos: {(_txt(r.get('Observação')) or 'sem tipo').lower()} que já existe como pendente (não duplica)"] += 1
                continue
            conta = _txt(r.get("Conta")) or ("Contabilizei Bank" if escopo == "PJ" else _conta_pf(desc))
            origem = _txt(r.get("ID")) or f"{aba}:{linha}"
            if obs == "eventual":
                venc = r.get("Vencimento")
                data = venc.date().isoformat() if isinstance(venc, dt.datetime) else vencimento(dia, INICIO)
                p["lancamentos"].append(dict.fromkeys(COLUNAS) | dict(
                    escopo=escopo, tipo="saida", data=data, valor=valor, categoria=_cat(r.get("Categoria")),
                    descricao=desc, conta=conta, observacao="custo eventual", origem=origem))
                continue
            if obs == "parcelado":   # a planilha não diz quantas parcelas faltam: fim fica em aberto
                p["revisar"].append(f"{origem}: parcelado sem pendentes futuros; recorrência criada sem fim, ajuste o fim")
            p["recorrencias"].append(dict(
                escopo=escopo, descricao=desc, categoria=_cat(r.get("Categoria")), valor=valor, dia_vencimento=dia,
                fixo=fixo, conta=conta, inicio=INICIO, fim=None, origem=origem))


def planejar(wb, clientes: dict) -> dict:
    """Lê a planilha e devolve o plano (nada é gravado). clientes: {nome normalizado: id}."""
    p = dict(lancamentos=[], recorrencias=[], lidas=Counter(), descartes=Counter(), arredondados=0,
             sem_cliente=Counter(), revisar=[])

    def pj(r, iso):
        tipo = _tipo(r)
        if not tipo:
            return None
        contra = _txt(r.get("Cliente/Fornecedor"))
        cid = clientes.get(_norm(contra)) if tipo == "entrada" else None
        if tipo == "entrada" and contra and not cid:
            p["sem_cliente"][_norm(contra)] += 1
        aulas = r.get("Aulas")
        return dict(escopo="PJ", tipo=tipo, pago_em=_pago(r, iso), categoria=_cat(r.get("Categoria")),
                    descricao=_txt(r.get("Descrição")), contraparte=contra, cliente_id=cid, conta="Contabilizei Bank",
                    forma=_forma(r.get("Forma Pagamento")),
                    nfse={"emitida": 1, "nao emitida": 0}.get(_norm(r.get("Emissão NFS-e"))),
                    observacao=f"aulas: {int(aulas)}" if isinstance(aulas, (int, float)) else None)

    def pf(r, iso):
        tipo = _tipo(r)
        if not tipo:
            return None
        desc = _txt(r.get("Descrição"))
        num, total = _parcela(r.get("Parcela"))
        return dict(escopo="PF", tipo=tipo, pago_em=_pago(r, iso), categoria=_cat(r.get("Categoria")), descricao=desc,
                    conta=_conta_pf(desc), forma=_forma(r.get("Forma Pagto")), parcela_num=num, parcela_total=total,
                    observacao=_obs(r.get("Forma"), r.get("Observação")))

    def cc(r, iso):
        cartao = _txt(r.get("Cartão")) or "sem nome"
        num, total = _parcela(r.get("Parcela"))
        return dict(escopo="PF", tipo="saida", pago_em=_pago(r, iso), categoria=_cat(r.get("Categoria") or "Fatura"),
                    descricao=f"Fatura {cartao}", conta=f"Cartão {cartao}", forma=_forma(r.get("Forma de Pagamento")),
                    parcela_num=num, parcela_total=total, observacao=_obs(r.get("Tipo"), r.get("Observação")))

    def ep(r, iso):   # entrada pontual: já recebida na competência
        if _norm(r.get("Incluir?")) != "sim":
            return None
        fonte = _txt(r.get("Fonte"))
        desc = _txt(r.get("Descrição"))
        return dict(escopo="PF", tipo="entrada", pago_em=iso, categoria="Salário" if "clt" in _norm(fonte) else "Outros",
                    descricao=desc, conta=_conta_pf(desc), observacao=fonte)

    _percorrer(wb, p, "Transações Empresariais", "Transações Empresariais", "tblTransaçõesEmpresariais", "PJ", pj)
    _percorrer(wb, p, "Transações Pessoais", "Transações Pessoais", "Tabela15", "PF", pf)
    _percorrer(wb, p, "Cartões", "Cartões", "tblCartoes", "CC", cc, col_valor="Valor")
    _percorrer(wb, p, "Entradas pontuais", "Painel Central", "tblEntradasPontuais", "EP", ep,
               col_data="Competência", col_valor="Valor")
    _custos(wb, p)   # depois das transações: precisa saber quais pendentes já existem
    # armadilha 10: pares iguais em data, tipo, valor e categoria só são listados, nunca removidos
    grupos = defaultdict(list)
    for l in p["lancamentos"]:
        if l["origem"][:3] in ("PJ:", "PF:"):
            grupos[(l["escopo"], l["data"], l["tipo"], l["valor"], l["categoria"])].append(l)
    p["suspeitos"] = [g for g in grupos.values() if len(g) > 1]
    return p


def _clientes(con) -> dict:
    return {_norm(r["nome"]): r["id"] for r in con.execute("SELECT id, nome FROM clientes WHERE excluido_em IS NULL")}


def _existentes(con, tabela: str) -> set:
    try:
        return {r[0] for r in con.execute(f"SELECT origem FROM {tabela} WHERE origem IS NOT NULL")}
    except sqlite3.OperationalError:   # banco real ainda sem as tabelas novas (a simulação não cria nada)
        return set()


def gravar(con, p) -> tuple:
    """INSERT OR IGNORE pela origem. Devolve (lançamentos novos, recorrências novas)."""
    with con:
        cur_r = con.executemany(
            f"INSERT OR IGNORE INTO recorrencias ({', '.join(COLUNAS_REC)}) VALUES ({', '.join(':' + c for c in COLUNAS_REC)})",
            p["recorrencias"])
        cur_l = con.executemany(
            f"INSERT OR IGNORE INTO lancamentos ({', '.join(COLUNAS)}) VALUES ({', '.join(':' + c for c in COLUNAS)})",
            p["lancamentos"])
    return cur_l.rowcount, cur_r.rowcount


def relatorio(p, con) -> str:
    ex_l, ex_r = _existentes(con, "lancamentos"), _existentes(con, "recorrencias")
    lan, rec = p["lancamentos"], p["recorrencias"]
    out = []
    out.append("Linhas lidas (sem as totalmente vazias): " + ", ".join(f"{k} {v}" for k, v in p["lidas"].items()))
    out.append("")
    out.append("Aproveitadas por origem / escopo / tipo:")
    for (orig, esc, tipo), n in sorted(Counter((l["origem"].split(":")[0] if ":" in l["origem"] else "custos",
                                                l["escopo"], l["tipo"]) for l in lan).items()):
        out.append(f"  {orig:7} {esc} {tipo:7} {n}")
    novos_l = sum(l["origem"] not in ex_l for l in lan)
    novos_r = sum(r["origem"] not in ex_r for r in rec)
    out.append("")
    out.append(f"Lançamentos: {len(lan)} no total, {novos_l} novos, {len(lan) - novos_l} já no CRM")
    out.append(f"Recorrências: {len(rec)} no total, {novos_r} novas, {len(rec) - novos_r} já no CRM")
    out.append(f"Valores arredondados para 2 casas: {int(p['arredondados'])}")
    out.append("")
    out.append("Descartes por motivo:")
    out += [f"  {n:4}  {m}" for m, n in sorted(p["descartes"].items())] or ["  nenhum"]
    if p["revisar"]:
        out += ["", "Para revisar:"] + [f"  {t}" for t in p["revisar"]]
    out.append("")
    out.append(f"Possíveis duplicidades (só listadas, nada é removido): {len(p['suspeitos'])}")
    for g in p["suspeitos"]:
        a = g[0]
        igual = "descrição igual" if len({_norm(l["descricao"]) for l in g}) == 1 else "descrição diferente"
        out.append(f"  {' + '.join(l['origem'] for l in g)}: {a['data']} {a['tipo']} {a['valor']:.2f} {a['categoria']} ({igual})")
    out.append("")
    out.append(f"Entradas PJ sem aluno correspondente: {len(p['sem_cliente'])} contrapartes, "
               f"{sum(p['sem_cliente'].values())} lançamentos (ficam sem cliente_id, nome em contraparte)")
    fora = Counter()
    for itens in (lan, rec):
        for x in itens:
            sufixo = "pj" if x["escopo"] == "PJ" else "pf"
            if x["categoria"] not in db.LISTAS_INICIAIS[f"lista_categorias_{sufixo}"]:
                fora[f"categoria {x['escopo']}: {x['categoria']}"] += 1
            if x.get("conta") and x["conta"] not in db.LISTAS_INICIAIS["lista_contas"]:
                fora[f"conta: {x['conta']}"] += 1
    for l in lan:
        if l["forma"] and l["forma"] not in db.LISTAS_INICIAIS["lista_formas_lancamento"]:
            fora[f"forma: {l['forma']}"] += 1
    out.append("Fora das listas de Configurações: " + (", ".join(f"{k} ({n})" for k, n in sorted(fora.items())) or "nada"))
    meses = defaultdict(lambda: Counter())
    for l in lan:
        if l["pago_em"]:
            meses[l["data"][:7]][(l["escopo"], l["tipo"])] += l["valor"]
    out.append("")
    out.append("Totais pagos por mês (para conferir com a planilha):")
    out.append("  mês       PJ entradas   PJ saídas  PF entradas   PF saídas")
    for m in sorted(meses):
        c = meses[m]
        out.append(f"  {m}  " + "  ".join(f"{c[k]:>11.2f}" for k in (("PJ", "entrada"), ("PJ", "saida"), ("PF", "entrada"), ("PF", "saida"))))
    return "\n".join(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Migra a planilha Contabilidade.xlsx para o financeiro do CRM.")
    ap.add_argument("--gravar", action="store_true", help="grava no banco (padrão: só simula)")
    ap.add_argument("--arquivo", type=Path, default=db.PLANILHA)
    args = ap.parse_args(argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")   # avisos de extensões do Excel que o openpyxl ignora
            # read_only=True não oferece ws.tables, que o importador também usa; nunca chamamos wb.save()
            wb = openpyxl.load_workbook(args.arquivo, data_only=True)
    except PermissionError:
        print(f"Não consegui abrir {args.arquivo}: o arquivo está aberto no Excel ou travado pelo OneDrive. "
              "Feche o Excel e rode de novo.")
        return 1
    except FileNotFoundError:
        print(f"Planilha não encontrada: {args.arquivo}")
        return 1
    try:
        if args.gravar:
            if db.DB_PATH.exists():
                print("Backup:", db.fazer_backup("antes-financeiro"))
            db.iniciar_banco()   # cria as tabelas novas no banco real, se ainda não existirem
            con = db.conectar()
        elif db.DB_PATH.exists():   # simulação: banco só para leitura
            con = sqlite3.connect(db.DB_PATH.as_uri() + "?mode=ro", uri=True)
            con.row_factory = sqlite3.Row
        else:
            con = sqlite3.connect(":memory:")
            con.row_factory = sqlite3.Row
            con.execute("CREATE TABLE clientes (id INTEGER PRIMARY KEY, nome TEXT, excluido_em TEXT)")
        try:
            plano = planejar(wb, _clientes(con))
            print(relatorio(plano, con))   # antes de gravar, para mostrar o que é novo
            if args.gravar:
                novos_l, novos_r = gravar(con, plano)
                print(f"\nGRAVADO: {novos_l} lançamentos e {novos_r} recorrências novos")
            else:
                print("\nSIMULAÇÃO: nada foi gravado. Para gravar: python -m app.migrar_financeiro --gravar")
        finally:
            con.close()
    finally:
        wb.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
