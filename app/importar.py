"""Importa dados da planilha Contabilidade.xlsx (somente leitura) sem duplicar registros."""
import json
import sqlite3
from pathlib import Path

import openpyxl

# ---- normalização de nomes da planilha -> nomes do CRM ----
CANAIS = {
    "CLT": "Outro",
    "Totem-CIA": "Parceria (síndico/academia)",
    "Conecta": "Conecta Fitness",
    "Instagram": "Instagram orgânico",
    "PersonalGo": "PersonalGo",
}
TIPOS = {
    "Pacotes de Aulas": "Pacote de Aulas",
    "Pacotes": "Pacote de Aulas",
    "Mensalidade (Aulas Presenciais)": "Mensalidade",
    "Mensalidades": "Mensalidade",
    "Aulas avulsas": "Aula avulsa",
}
LOCAIS = {"CIA Athletica": "CIA"}
DIAS = {"Segunda": 1, "Terça": 2, "Quarta": 3, "Quinta": 4, "Sexta": 5, "Sábado": 6, "Domingo": 7}
STATUS_SLOT = {"Disponível": "livre", "Ocupado": "ocupado", "Indisponível": "indisponivel"}


def _txt(v):
    return str(v).strip() if v not in (None, "") else None


def _num(v, padrao=0.0):
    try:
        return float(v)
    except (TypeError, ValueError):
        return padrao


def canal(v, canais_validos):
    v = _txt(v)
    if not v:
        return None
    v = CANAIS.get(v, v)
    return v if v in canais_validos else "Outro"


def tipo(v):
    v = _txt(v)
    return TIPOS.get(v, v) if v else None


def local(v):
    v = _txt(v)
    return LOCAIS.get(v, v) if v else None


def ler_tabela(wb, aba: str, tabela: str) -> list[dict]:
    """Linhas de uma Tabela do Excel como dicionários (cabeçalho -> valor). Ignora linhas vazias."""
    ws = wb[aba]
    ref = ws.tables[tabela].ref
    linhas = ws[ref]
    cab = [str(c.value).strip() for c in linhas[0]]
    out = []
    for linha in linhas[1:]:
        vals = [c.value for c in linha]
        if all(v in (None, "") for v in vals):
            continue
        out.append(dict(zip(cab, vals)))
    return out


# ---- cada função devolve {"inseridos": n, "atualizados": n} ----
def importar_clientes(con: sqlite3.Connection, wb, canais_validos) -> dict:
    ins = upd = 0
    for r in ler_tabela(wb, "Clientes", "Tabela16"):
        if not _txt(r.get("Nome")) or r.get("ID") in (None, ""):
            continue
        t = tipo(r.get("Tipo Servico"))
        aulas_mes = int(_num(r.get("Aulas/Planos/Mês")))
        freq = int(_num(r.get("Freq. Semanal")))
        if t == "Pacote de Aulas" and aulas_mes:
            plano = f"Pacote {aulas_mes} aulas"
        elif t == "Mensalidade" and freq:
            plano = f"Mensalidade {freq}x/semana"
        else:
            plano = t
        dados = dict(
            planilha_id=int(r["ID"]),
            nome=_txt(r["Nome"]),
            valor=_num(r.get("Receita Mensal")) or _num(r.get("Valor Aula/Plano")),
            local=local(r.get("Local")),
            tipo_servico=t,
            plano=plano,
            forma_pagamento=_txt(r.get("Forma Pgto.")),
            status=(_txt(r.get("Status")) or "ativo").lower(),
            canal_origem=canal(r.get("Canal de Origem"), canais_validos),
        )
        existe = con.execute("SELECT id FROM clientes WHERE planilha_id=?", (dados["planilha_id"],)).fetchone()
        if existe:
            con.execute(
                """UPDATE clientes SET nome=:nome, valor=:valor, local=:local, tipo_servico=:tipo_servico,
                   plano=:plano, forma_pagamento=:forma_pagamento, status=:status, canal_origem=:canal_origem,
                   atualizado_em=datetime('now','localtime') WHERE planilha_id=:planilha_id""",
                dados,
            )
            upd += 1
        else:
            con.execute(
                """INSERT INTO clientes (planilha_id, nome, valor, local, tipo_servico, plano, forma_pagamento,
                   status, canal_origem) VALUES (:planilha_id, :nome, :valor, :local, :tipo_servico, :plano,
                   :forma_pagamento, :status, :canal_origem)""",
                dados,
            )
            ins += 1
    return {"inseridos": ins, "atualizados": upd}


def _upsert_preco(con, dados, existentes: set) -> bool:
    """Insere ou atualiza pela chave composta. Devolve True se era novo."""
    chave = (dados["tipo_servico"], dados["local"], dados["forma_pagamento"], dados["qtd_aulas"], dados["freq_semanal"])
    con.execute(
        """INSERT INTO precos_base (tipo_servico, local, forma_pagamento, qtd_aulas, freq_semanal, sessoes_mes,
           preco_base, custo_operacional, mao_obra, custo_total, origem)
           VALUES (:tipo_servico, :local, :forma_pagamento, :qtd_aulas, :freq_semanal, :sessoes_mes,
           :preco_base, :custo_operacional, :mao_obra, :custo_total, :origem)
           ON CONFLICT(tipo_servico, local, forma_pagamento, qtd_aulas, freq_semanal) DO UPDATE SET
           sessoes_mes=excluded.sessoes_mes, preco_base=excluded.preco_base,
           custo_operacional=excluded.custo_operacional, mao_obra=excluded.mao_obra,
           custo_total=excluded.custo_total, origem=excluded.origem""",
        dados,
    )
    novo = chave not in existentes
    existentes.add(chave)
    return novo


def importar_precos(con: sqlite3.Connection, wb) -> dict:
    ins = upd = 0
    existentes = {tuple(r) for r in con.execute(
        "SELECT tipo_servico, local, forma_pagamento, qtd_aulas, freq_semanal FROM precos_base")}
    for r in ler_tabela(wb, "Comparador", "tblComparadorPrecos"):
        if not _txt(r.get("Tipo Serviço")) or r.get("Preço Sugerido") in (None, ""):
            continue
        t = tipo(r["Tipo Serviço"])
        qtd = int(_num(r.get("Qtd Aulas/Planos")))
        dados = dict(
            tipo_servico=t, local=local(r.get("Local")) or "Online",
            forma_pagamento=_txt(r.get("Forma Pagamento")) or "PIX",
            qtd_aulas=qtd if t == "Pacote de Aulas" else 0, freq_semanal=0,
            sessoes_mes=qtd if t in ("Pacote de Aulas", "Aula avulsa") else (1 if t == "Consultoria Híbrida" else 0),
            preco_base=_num(r["Preço Sugerido"]), custo_operacional=_num(r.get("Custo Operacional")),
            mao_obra=_num(r.get("Mão-de-Obras")), custo_total=_num(r.get("Custo Total")), origem="comparador",
        )
        if _upsert_preco(con, dados, existentes):
            ins += 1
        else:
            upd += 1
    for r in ler_tabela(wb, "Comparador", "Tabela12"):
        if r.get("Freq. Semanal") in (None, "") or r.get("Preço Sugerido") in (None, ""):
            continue
        dados = dict(
            tipo_servico="Mensalidade", local=local(r.get("Local")) or "Online", forma_pagamento="PIX",
            qtd_aulas=0, freq_semanal=int(_num(r["Freq. Semanal"])), sessoes_mes=_num(r.get("Aulas Nominais")),
            preco_base=_num(r["Preço Sugerido"]), custo_operacional=_num(r.get("Custo Operacional")),
            mao_obra=_num(r.get("Mão-de-Obra")), custo_total=_num(r.get("Custo Total")), origem="tabela12",
        )
        if _upsert_preco(con, dados, existentes):
            ins += 1
        else:
            upd += 1
    return {"inseridos": ins, "atualizados": upd}


def importar_matriz(con: sqlite3.Connection, wb) -> dict:
    ins = upd = 0
    for r in ler_tabela(wb, "Matriz de Horarios", "tblMatrizHorariosNormalizada"):
        dia = DIAS.get(_txt(r.get("Dia")))
        if not dia or r.get("Horário") in (None, ""):
            continue
        tc = "Mensalidades" if "ensalidade" in str(r.get("Tipo cálculo")) else "Pacotes"
        h = r["Horário"]
        hora = h.hour if hasattr(h, "hour") else int(round(_num(h) * 24))   # célula pode vir como time ou fração
        mult = r.get("Modificador Preço")
        dados = dict(
            tipo_calculo=tc, dia=dia, local=local(r.get("Local")), hora=hora,
            status=STATUS_SLOT.get(_txt(r.get("Status disponibilidade")), "livre"),
            cliente_nome=_txt(r.get("Cliente/Observação")), bloqueio=_txt(r.get("Bloqueio fixo")),
            multiplicador=_num(mult, 1.0) if mult not in (None, "") else 1.0,
        )
        existe = con.execute(
            "SELECT id, status, cliente_id FROM grade_slots WHERE tipo_calculo=? AND dia=? AND local=? AND hora=?",
            (tc, dia, dados["local"], hora),
        ).fetchone()
        if existe:
            if existe["status"] == "ocupado" and existe["cliente_id"]:
                # célula ocupada por cliente vinculado no CRM: só o multiplicador vem da planilha
                con.execute("UPDATE grade_slots SET multiplicador=? WHERE id=?", (dados["multiplicador"], existe["id"]))
            else:
                con.execute(
                    """UPDATE grade_slots SET status=:status, cliente_nome=:cliente_nome, bloqueio=:bloqueio,
                       multiplicador=:multiplicador WHERE id=:id""", {**dados, "id": existe["id"]})
            upd += 1
        else:
            con.execute(
                """INSERT INTO grade_slots (tipo_calculo, dia, local, hora, status, cliente_nome, bloqueio, multiplicador)
                   VALUES (:tipo_calculo, :dia, :local, :hora, :status, :cliente_nome, :bloqueio, :multiplicador)""",
                dados,
            )
            ins += 1
    return {"inseridos": ins, "atualizados": upd}


def importar_marketing(con: sqlite3.Connection, wb, canais_validos) -> dict:
    ins = upd = 0
    for r in ler_tabela(wb, "Marketing", "tblMarketingMarketplaces"):
        mes, app = _txt(r.get("Período")), _txt(r.get("Aplicativo"))
        if not mes or not app:
            continue
        dados = dict(
            mes=mes[:7], canal=canal(app, canais_validos), mensalidade=_num(r.get("Mensalidade")),
            anuncios=_num(r.get("Anúncios")), comissoes=_num(r.get("Comissões")), outros=_num(r.get("Outras Taxas")),
            leads=int(_num(r.get("Leads"))), agendamentos=int(_num(r.get("Agendamentos"))),
            clientes=int(_num(r.get("Clientes"))), observacao=_txt(r.get("Observação")),
        )
        existe = con.execute("SELECT id FROM gastos_marketing WHERE mes=? AND canal=?", (dados["mes"], dados["canal"])).fetchone()
        if existe:
            con.execute(
                """UPDATE gastos_marketing SET mensalidade=:mensalidade, anuncios=:anuncios, comissoes=:comissoes,
                   outros=:outros, leads=:leads, agendamentos=:agendamentos, clientes=:clientes,
                   observacao=:observacao WHERE mes=:mes AND canal=:canal""", dados)
            upd += 1
        else:
            con.execute(
                """INSERT INTO gastos_marketing (mes, canal, mensalidade, anuncios, comissoes, outros, leads,
                   agendamentos, clientes, observacao) VALUES (:mes, :canal, :mensalidade, :anuncios, :comissoes,
                   :outros, :leads, :agendamentos, :clientes, :observacao)""", dados)
            ins += 1
    return {"inseridos": ins, "atualizados": upd}


def importar_config(con: sqlite3.Connection, wb) -> dict:
    """Parâmetros da planilha entram só se a chave ainda não existe (não desfaz edições da tela Configurações)."""
    novos = []
    for r in ler_tabela(wb, "Comparador", "Tabela20"):           # margem mínima por serviço
        if _txt(r.get("Tipo Serviço")):
            novos.append(("margem_minima", tipo(r["Tipo Serviço"]), _num(r.get("Margem")), None, 9))
    for r in ler_tabela(wb, "Comparador", "Tabela13"):           # taxas de pagamento
        if _txt(r.get("Forma Pgto")):
            novos.append(("taxa_pagamento", _txt(r["Forma Pgto"]), _num(r.get("Taxa Total")), None, 9))
    for r in ler_tabela(wb, "Comparador", "Tabela18"):           # taxa CIA por pessoa
        nome = _txt(r.get("Taxas")) or ""
        if nome.startswith("CIA") and r.get("Nº de pessoas") not in (None, "", "N/A"):
            novos.append(("taxa_cia", str(int(_num(r["Nº de pessoas"]))), _num(r.get("Valor [R$]")), None, 9))
    for r in ler_tabela(wb, "Config_Servicos", "tblPrecificacaoGrupos"):
        p = r.get("Pessoas mínimas")
        if p not in (None, ""):
            novos.append(("grupo_multiplicador", str(int(_num(p))), _num(r.get("Multiplicador total"), 1), None, 9))
            novos.append(("grupo_mao_obra_extra", str(int(_num(p))), _num(r.get("Esforço extra MO")), None, 9))
    for r in ler_tabela(wb, "Config_Servicos", "Tabela24"):      # mão de obra por aula conforme qtd
        if _txt(r.get("Tipo de Serviço")):
            novos.append(("mao_obra_aula", f"{tipo(r['Tipo de Serviço'])}|{int(_num(r.get('Qtd Aulas')))}",
                          _num(r.get("Mão-de-Obra")), None, 9))
    antes = con.execute("SELECT COUNT(*) FROM parametros").fetchone()[0]
    con.executemany("INSERT OR IGNORE INTO parametros (grupo, chave, valor, valor_texto, ordem) VALUES (?,?,?,?,?)", novos)
    depois = con.execute("SELECT COUNT(*) FROM parametros").fetchone()[0]
    return {"inseridos": depois - antes, "atualizados": 0}


def importar_tudo(con: sqlite3.Connection, caminho: Path) -> dict:
    """Roda todas as importações numa transação e registra o resumo em `importacoes`."""
    wb = openpyxl.load_workbook(caminho, data_only=True)   # nunca chamamos wb.save()
    try:
        canais_validos = [r[0] for r in con.execute("SELECT chave FROM parametros WHERE grupo='lista_canais'")]
        with con:
            resumo = {
                "Clientes": importar_clientes(con, wb, canais_validos),
                "Preços": importar_precos(con, wb),
                "Grade de horários": importar_matriz(con, wb),
                "Marketing": importar_marketing(con, wb, canais_validos),
                "Configurações": importar_config(con, wb),
            }
            con.execute("INSERT INTO importacoes (arquivo, resumo) VALUES (?, ?)",
                        (str(caminho), json.dumps(resumo, ensure_ascii=False)))
    finally:
        wb.close()
    return resumo
