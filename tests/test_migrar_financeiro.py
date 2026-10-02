"""Migração da planilha: xlsx sintético, descartes, arredondamento e idempotência do --gravar."""
import datetime as dt

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from app import db, migrar_financeiro


@pytest.fixture
def con_temp(tmp_path, monkeypatch):
    db_path = tmp_path / "crm.db"
    monkeypatch.setattr(db, "DB_PATH", db_path)
    monkeypatch.setattr(db, "BACKUP_DIR", tmp_path / "backups")
    db.iniciar_banco(db_path)
    con = db.conectar(db_path)
    yield con
    con.close()


def _tabela(wb, aba, nome, cab, linhas, col=2, linha_cab=2):
    """Escreve cabeçalho + linhas (None = linha em branco) e registra como Tabela do Excel."""
    ws = wb[aba] if aba in wb.sheetnames else wb.create_sheet(aba)
    for j, c in enumerate(cab):
        ws.cell(linha_cab, col + j, c)
    for i, vals in enumerate(linhas, 1):
        for j, v in enumerate(vals or []):
            ws.cell(linha_cab + i, col + j, v)
    ref = f"{get_column_letter(col)}{linha_cab}:{get_column_letter(col + len(cab) - 1)}{linha_cab + len(linhas)}"
    ws.add_table(Table(displayName=nome, ref=ref))


def _planilha(caminho):
    d = lambda dia, mes=9: dt.datetime(2026, mes, dia)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cab_custo = ["Descrição", "Categoria", "Valor", "Dia Venc.", "Vencimento", "Baixa manual?", "Competência baixa",
                 "Data pagamento", "Status consolidado", "Conta", "Observação", "ID", "Conciliação"]
    _tabela(wb, "Transações Empresariais", "tblTransaçõesEmpresariais",
            ["Data", "Tipo", "Cliente/Fornecedor", "Categoria", "Valor (R$)", "Forma Pagamento", "Status", "Descrição",
             "Mês", "Ano", "Emissão NFS-e", "Aulas"],
            [[d(1), "Entrada", "joao  ALUNO", "Pacotes de Aulas", 10.123, "PIX", "Pago", None, "setembro", 2026, "Emitida", 4],
             [d(2), "Saída", "Fornecedor X", "Aluguel", 50, "PIX", "Pendente", "Sala", "setembro", 2026, None, None],
             None,
             [None, None, None, None, 999, None, None, None, None, None, None, None]])   # total: sem data
    _tabela(wb, "Transações Pessoais", "Tabela15",
            ["Mês", "Data", "Tipo", "Categoria", "Descrição", "Valor (R$)", "Forma Pagto", "Status", "Forma", "Parcela",
             "Observação", "ID Obrigação", "Conciliação custo"],
            [[None, d(3), "Saída", "Outro", "Mercado Pago lanche", 20.5, "Débito", "Pago", "Eventual", "2 de 12", None, None, None],
             [None, d(20, 11), "Saída", "Moradia", "Aluguel casa", 1000, "Boleto", "Pendente", "Mensal", None, None, None, None],
             [None, "Total", None, None, None, 5000, None, None, None, None, None, None, None]])   # total em texto
    _tabela(wb, "Cartões", "tblCartoes",
            ["Data", "Cartão", "Descrição", "Categoria", "Valor", "Forma de Pagamento", "Status", "Tipo", "Parcela",
             "Observação", "Mês"],
            [[d(5), "Inter", None, "Fatura", 300, "PIX", "Pago", "Variável", None, None, "setembro"],
             [None, None, None, None, None, None, None, None, None, None, "Janeiro"]])   # linha fantasma
    _tabela(wb, "Custos Empresariais", "tblCustosEmpresaFixos", cab_custo,
            [["Software X", "Software/Assinaturas", 30, 31, d(30), "Não", None, d(15), "Pendente", "Contabilizei Bank", "Recorrente", "CEF-001", None],
             None], col=2)
    _tabela(wb, "Custos Empresariais", "tblCustosEmpresaVariaveis", cab_custo, [None], col=16)
    _tabela(wb, "Custos Pessoais", "tblCustosPessoaisFixos", cab_custo,
            [["Aluguel casa", "Moradia", 1000, 20, d(20), "Não", None, None, "Pendente", "Inter", "Recorrente", "CPF-001", None]], col=2)
    _tabela(wb, "Custos Pessoais", "tblCustosPessoaisVariaveis", cab_custo,
            [["Conserto", "Pessoal", 80, 25, d(25), "Não", None, None, "Pendente", "Inter", "Eventual", "CPV-001", None]], col=16)
    _tabela(wb, "Painel Central", "tblEntradasPontuais", ["Competência", "Fonte", "Descrição", "Valor", "Incluir?"],
            [[d(1), "Renda CLT / Rescisão", "Rescisão", 700, "Sim"], [d(1), "Ajustes manuais", "Ignorar", 5, "Não"]])
    wb.save(caminho)


def _contagens(con):
    return [con.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in ("lancamentos", "recorrencias")]


def test_planejar_descarta_arredonda_e_nao_duplica(tmp_path, con_temp):
    xlsx = tmp_path / "Contabilidade.xlsx"
    _planilha(xlsx)
    con_temp.execute("INSERT INTO clientes (nome) VALUES ('João Aluno')")
    con_temp.commit()

    # simulação não grava nada
    assert migrar_financeiro.main(["--arquivo", str(xlsx)]) == 0
    assert _contagens(con_temp) == [0, 0]

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    p = migrar_financeiro.planejar(wb, migrar_financeiro._clientes(con_temp))
    assert p["descartes"]["Transações Empresariais: sem data (total ou linha fantasma)"] == 1   # total PJ
    assert p["descartes"]["Transações Pessoais: data em texto (linha de total)"] == 1
    assert p["descartes"]["Cartões: sem data (total ou linha fantasma)"] == 1                   # fantasma
    assert p["descartes"]["Custos: 'Data pagamento' não vira lançamento (já está nas transações)"] == 1
    assert p["descartes"]["Custos: recorrente que já existe como pendente (não duplica)"] == 1    # CPF-001 já é pendente PF
    assert p["arredondados"] == 1
    por_origem = {l["origem"]: l for l in p["lancamentos"]}
    assert por_origem["PJ:3"]["valor"] == 10.12 and por_origem["PJ:3"]["cliente_id"] == 1 and por_origem["PJ:3"]["nfse"] == 1
    assert por_origem["PJ:4"]["pago_em"] is None
    assert por_origem["PF:3"]["conta"] == "Mercado Pago" and por_origem["PF:3"]["forma"] == "Cartão Débito"
    assert (por_origem["PF:3"]["parcela_num"], por_origem["PF:3"]["parcela_total"]) == (2, 12)
    assert por_origem["PF:3"]["categoria"] == "Outros"
    assert por_origem["CC:3"]["conta"] == "Cartão Inter" and por_origem["CC:3"]["descricao"] == "Fatura Inter"
    assert por_origem["EP:3"]["categoria"] == "Salário" and "EP:4" not in por_origem
    assert por_origem["CPV-001"]["pago_em"] is None and por_origem["CPV-001"]["data"] == "2026-09-25"
    assert [r["origem"] for r in p["recorrencias"]] == ["CEF-001"]

    # --gravar duas vezes: backup antes, e a segunda não duplica nada
    assert migrar_financeiro.main(["--arquivo", str(xlsx), "--gravar"]) == 0
    primeira = _contagens(con_temp)
    assert primeira == [len(p["lancamentos"]), 1]
    assert list((db.BACKUP_DIR).glob("antes-financeiro-*.db"))
    assert migrar_financeiro.main(["--arquivo", str(xlsx), "--gravar"]) == 0
    assert _contagens(con_temp) == primeira


def test_planilha_aberta_no_excel(tmp_path, monkeypatch, capsys):
    def trava(*a, **k):
        raise PermissionError

    monkeypatch.setattr(migrar_financeiro.openpyxl, "load_workbook", trava)
    assert migrar_financeiro.main(["--arquivo", str(tmp_path / "x.xlsx")]) == 1
    assert "Feche o Excel" in capsys.readouterr().out
